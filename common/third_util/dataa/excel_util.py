from openpyxl import load_workbook, Workbook
from openpyxl.utils.dataframe import dataframe_to_rows
from openpyxl.cell.cell import MergedCell
from typing import Optional, Union, Dict, List
import os
import pandas as pd


class ExcelUtil:
    SUPPORTED_FORMATS = ['.xlsx', '.xls', '.xlsm', '.csv']

    def __init__(self, path: Optional[str] = None):
        self._path = path
        self._data: Dict[str, pd.DataFrame] = {}
        self._workbook: Optional[Workbook] = None

    def load(self, path: Optional[str] = None) -> "ExcelUtil":
        self._path = path or self._path
        if not self._path:
            raise ValueError("path is required")
        if not os.path.exists(self._path):
            raise FileNotFoundError(f"file not found: {self._path}")
        self._data = self.read_all_sheets()
        return self

    def read_sheet(self, sheet_name: Optional[str] = None) -> pd.DataFrame:
        if not self._path:
            raise ValueError("path is required")
        ext = os.path.splitext(self._path)[1].lower()
        if ext == '.csv':
            return pd.read_csv(self._path)
        sheet_name = sheet_name or 0
        return pd.read_excel(self._path, sheet_name=sheet_name)

    def read_all_sheets(self) -> Dict[str, pd.DataFrame]:
        if not self._path:
            raise ValueError("path is required")
        ext = os.path.splitext(self._path)[1].lower()
        if ext == '.csv':
            name = os.path.splitext(os.path.basename(self._path))[0]
            return {name: pd.read_csv(self._path)}
        return pd.read_excel(self._path, sheet_name=None)

    def read_range(
        self,
        sheet_name: str,
        start_row: int,
        end_row: int,
        start_col: int,
        end_col: int,
    ) -> pd.DataFrame:
        df = self.read_sheet(sheet_name)
        return df.iloc[start_row:end_row, start_col:end_col]

    def get_sheet(self, sheet_name: str) -> pd.DataFrame:
        if sheet_name not in self._data:
            raise KeyError(f"sheet not found: {sheet_name}")
        return self._data[sheet_name]

    def get_sheet_names(self) -> List[str]:
        if self._path and os.path.exists(self._path):
            ext = os.path.splitext(self._path)[1].lower()
            if ext == '.csv':
                return [os.path.splitext(os.path.basename(self._path))[0]]
            wb = load_workbook(self._path, read_only=True)
            return wb.sheetnames
        return list(self._data.keys())

    def get_row_count(self, sheet_name: Optional[str] = None) -> int:
        df = self.get_sheet(sheet_name) if sheet_name else list(self._data.values())[0]
        return len(df)

    def get_col_count(self, sheet_name: Optional[str] = None) -> int:
        df = self.get_sheet(sheet_name) if sheet_name else list(self._data.values())[0]
        return len(df.columns)

    def write(self, path: str, format: Optional[str] = None) -> "ExcelUtil":
        ext = format or os.path.splitext(path)[1].lower()
        if ext == '.csv':
            first_sheet = list(self._data.keys())[0]
            self._data[first_sheet].to_csv(path, index=False)
            return self
        if ext in ['.xlsx', '.xlsm']:
            wb = Workbook()
            wb.remove(wb.active)
            for sheet_name, df in self._data.items():
                ws = wb.create_sheet(title=sheet_name)
                for r in dataframe_to_rows(df, index=False, header=True):
                    ws.append(r)
            wb.save(path)
            return self
        raise ValueError(f"unsupported format: {ext}")

    def write_sheet(
        self, sheet_name: str, data: Union[pd.DataFrame, List, Dict]
    ) -> "ExcelUtil":
        if isinstance(data, pd.DataFrame):
            self._data[sheet_name] = data
        elif isinstance(data, list):
            self._data[sheet_name] = pd.DataFrame(data)
        elif isinstance(data, dict):
            self._data[sheet_name] = pd.DataFrame.from_dict(data)
        return self

    def append_row(self, sheet_name: str, row: Union[List, Dict]) -> "ExcelUtil":
        if sheet_name not in self._data:
            raise KeyError(f"sheet not found: {sheet_name}")
        df = self._data[sheet_name]
        if isinstance(row, dict):
            self._data[sheet_name] = pd.concat([df, pd.DataFrame([row])], ignore_index=True)
        else:
            self._data[sheet_name] = pd.concat(
                [df, pd.DataFrame([row], columns=df.columns)], ignore_index=True
            )
        return self

    def save(self, path: Optional[str] = None) -> "ExcelUtil":
        output_path = path or self._path
        if not output_path:
            raise ValueError("path is required")
        return self.write(output_path)

    def to_json(self) -> Dict:
        result = {}
        for sheet_name, df in self._data.items():
            result[sheet_name] = df.to_dict(orient='records')
        return result

    def to_dict(self, sheet_name: Optional[str] = None) -> Union[Dict, List]:
        if sheet_name:
            return self._data[sheet_name].to_dict(orient='records')
        return self.to_json()

    def to_csv(self, path: str, sheet_name: Optional[str] = None) -> "ExcelUtil":
        if sheet_name:
            self._data[sheet_name].to_csv(path, index=False)
        else:
            first_sheet = list(self._data.keys())[0]
            self._data[first_sheet].to_csv(path, index=False)
        return self

    def copy_and_modify(
        self,
        source_path: str,
        output_path: Optional[str] = None,
        modifications: Optional[Dict[str, List[Dict]]] = None,
    ) -> "ExcelUtil":
        source_path = os.path.abspath(source_path)
        if not os.path.exists(source_path):
            raise FileNotFoundError(f"file not found: {source_path}")

        output_path = output_path or self._generate_output_path(source_path)
        wb = load_workbook(source_path)
        if modifications:
            for sheet_name, mods in modifications.items():
                if sheet_name not in wb.sheetnames:
                    raise KeyError(f"sheet not found: {sheet_name}")
                ws = wb[sheet_name]
                col_map = self._build_col_map(ws)
                for mod in mods:
                    row = mod.get("row")
                    col_name = mod.get("col")
                    value = mod.get("value")
                    if row is None or col_name is None:
                        raise ValueError(f"modification must have 'row' and 'col': {mod}")
                    if col_name not in col_map:
                        raise KeyError(f"column not found: {col_name} in sheet {sheet_name}")
                    col_idx = col_map[col_name]
                    cell = ws.cell(row=row, column=col_idx)
                    if isinstance(cell, MergedCell):
                        parent_cell = self._find_merged_parent(ws, row, col_idx)
                        if parent_cell:
                            parent_cell.value = value
                    else:
                        cell.value = value

        wb.save(output_path)
        self.load(output_path)
        return self

    def _generate_output_path(self, source_path: str) -> str:
        dir_path = os.path.dirname(source_path)
        filename = os.path.basename(source_path)
        name, ext = os.path.splitext(filename)
        new_name = f"{name}_modified{ext}"
        return os.path.join(dir_path, new_name)

    def _build_col_map(self, worksheet) -> Dict[str, int]:
        col_map = {}
        for col_idx in range(1, worksheet.max_column + 1):
            cell_value = worksheet.cell(row=1, column=col_idx).value
            if cell_value:
                col_map[str(cell_value)] = col_idx
        return col_map

    def _find_merged_parent(self, worksheet, row: int, col: int):
        for merged_range in worksheet.merged_cells.ranges:
            if merged_range.min_row <= row <= merged_range.max_row and merged_range.min_col <= col <= merged_range.max_col:
                return worksheet.cell(row=merged_range.min_row, column=merged_range.min_col)
        return None

    @staticmethod
    def is_excel_file(file_path: str) -> bool:
        ext = os.path.splitext(file_path)[1].lower()
        return ext in ExcelUtil.SUPPORTED_FORMATS